"""Build the customer DOCX from the maintained SOP. Requires python-docx."""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/sop/spp-entra-id.md'

def inline(p, text):
    pattern = r'(`[^`]+`|\[[^\]]+\]\([^)]+\))'
    for item in re.split(pattern, text):
        if not item:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', item)
        if link:
            label, url = link.groups()
            if not url.startswith('http'):
                rel = (SOURCE.parent / url).resolve().relative_to(ROOT).as_posix()
                url = 'https://github.com/BrianHsing/One-Identity-Safeguard/blob/main/' + rel
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
            r = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '175A7A'); props.append(color)
            r.append(props)
            t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)
        else:
            run = p.add_run(item.strip('`') if item.startswith('`') else item)
            if item.startswith('`'):
                run.font.name = 'Consolas'
                run.font.size = Pt(10)

def table(doc, rows):
    count = len(rows[0])
    widths = [47, 123] if count == 2 else [48, 64, 58]
    t = doc.add_table(rows=0, cols=count)
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Mm(width)
    for idx, values in enumerate(rows):
        cells = t.add_row().cells
        for cell, width, value in zip(cells, widths, values):
            cell.width = Mm(width)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            if len(rows) <= 4 and idx < len(rows)-1:
                p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.1
            inline(p, value)
            for run in p.runs: run.font.size = Pt(10)
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EFF3' if idx == 0 else ('F6F8FA' if idx%2 == 0 else 'FFFFFF')); tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e = OxmlElement('w:'+side); e.set(qn('w:w'), '85'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tcpr.append(margins)
            if idx == 0:
                for run in p.runs: run.bold = True
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx == 0: trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def build(output):
    text = SOURCE.read_text(encoding='utf-8-sig')
    # The customer copy retains the operational procedure, not the internal comparison.
    text = re.sub(r'### 參考文章整合與差異.*?(?=## 相關文件)', '', text, flags=re.S)
    text = text.replace('- [實機截圖與驗證範圍](environment-evidence.md)\n', '')
    body, refs = text.split('## 相關文件', 1)
    reference_links = [line[2:] for line in refs.splitlines() if line.startswith('- ')]
    reference_links.append('[流程參考 Notion](https://zihshuo976.notion.site/Safeguard-Entra-ID-SAML-2-0-3226df2f41d38004afa6e97fb067730c)')
    text = body + '## 相關文件\n\n' + '；'.join(reference_links) + '\n\n原廠部分舊連結可能轉址；9.0 版本差異須依現場驗收。\n'
    doc = Document()
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(18); sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(20); sec.right_margin = Mm(20)
    sec.header_distance = Mm(8); sec.footer_distance = Mm(8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        s = doc.styles[name]
        s.font.name = 'Microsoft JhengHei'
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        s.font.color.rgb = RGBColor.from_string('111111')
    normal = doc.styles['Normal']; normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    for name, size in [('Title',22),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
        s = doc.styles[name]; s.font.size = Pt(size); s.font.bold = True
        s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].font.color.rgb = RGBColor.from_string('555555')
    hp = sec.header.paragraphs[0]
    hp.add_run('One Identity Safeguard  |  Microsoft Entra ID').font.size = Pt(9)
    fp = sec.footer.paragraphs[0]; fp.alignment = 2
    fp.add_run('2026-10-01  ·  版本 1.0     第 ').font.size = Pt(9)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
    fp.add_run(' 頁').font.size = Pt(9)
    lines = text.splitlines(); i = 0
    page_starts = ('第四階段',)
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line: continue
        if line.startswith('# '):
            doc.add_paragraph('SPP 整合 Microsoft Entra ID', 'Title')
            doc.add_paragraph('SAML 外部同盟操作指引\n版本 1.0  |  2026-10-01')
        elif line.startswith('#'):
            level = len(line)-len(line.lstrip('#')); title = line[level:].strip()
            p = doc.add_paragraph(title, 'Heading '+str(min(level-1,3)))
            if title.startswith(page_starts) or title in ('操作步驟', '施工欄位表', '注意事項'):
                p.paragraph_format.page_break_before = True
        elif line.startswith('|'):
            rows = [[v.strip() for v in line.strip('|').split('|')]]
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = lines[i].strip(); i += 1
                if re.fullmatch(r'[|\s:\-]+',row): continue
                rows.append([v.strip() for v in row.strip('|').split('|')])
            table(doc, rows)
        elif line.startswith('!['):
            m = re.match(r'!\[([^]]+)\]\(([^)]+)\)',line)
            p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
            width = 125 if 'spp-federation-config' in m[2] else 160
            shape = p.add_run().add_picture(str((SOURCE.parent / m[2]).resolve()), width=Mm(width))
            shape._inline.docPr.set('descr', m[1])
        else:
            p = doc.add_paragraph(style='Caption' if re.match(r'圖 \d', line) else None)
            if line.startswith('- '): line = line[2:]
            inline(p, line)
    doc.core_properties.title = 'SPP 整合 Microsoft Entra ID 操作指引'
    doc.core_properties.subject = 'SAML 外部同盟部署與驗收'
    doc.core_properties.author = ''
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--output', type=Path, default=ROOT/'docs/customer/SPP整合Microsoft Entra ID操作指引.docx')
    build(ap.parse_args().output)
